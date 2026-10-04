"""One standard resume design shared by the editable and portable exports."""
from datetime import datetime, timezone
from io import BytesIO
import hashlib
import json
import os
from pathlib import Path
from urllib.parse import urlparse
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Flowable, Paragraph, PageBreak

from resume_ir import ResumeIR, iter_lines

RENDER_VERSION = "resume-standard-v1"
_ROOT = Path(__file__).resolve().parents[1]
_CONFIG_PATH = _ROOT / "assets/resume-standard-v1.json"
CONFIG = json.loads(_CONFIG_PATH.read_text())
PAGE = tuple(mm / 25.4 * 72 for mm in CONFIG["page_mm"])
MARGINS = CONFIG["margins_pt"]
WIDTH = PAGE[0] - MARGINS["left"] - MARGINS["right"]
_FONT = None


def _system_font_pair():
    windows = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    candidates = [
            (Path("/System/Library/Fonts/Supplemental/Arial.ttf"), Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf")),
            (windows / "arial.ttf", windows / "arialbd.ttf"),
            (Path("/usr/share/fonts/truetype/msttcorefonts/Arial.ttf"), Path("/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf")),
    ]
    return next(((a,b) for a,b in candidates if a.is_file() and b.is_file()), None)


def _font_choice():
    global _FONT
    if _FONT is None:
        chosen = _system_font_pair()
        family = CONFIG["preferred_font"] if chosen else CONFIG["fallback_font"]
        regular, bold = chosen or (_ROOT / "assets/fonts/NotoSans-Regular.ttf", _ROOT / "assets/fonts/NotoSans-Bold.ttf")
        pdfmetrics.registerFont(TTFont("RolewardStandard", str(regular)))
        pdfmetrics.registerFont(TTFont("RolewardStandard-Bold", str(bold)))
        pdfmetrics.registerFontFamily("RolewardStandard", normal="RolewardStandard", bold="RolewardStandard-Bold")
        _FONT = {"font_family":family, "font_fallback_used":chosen is None}
    return _FONT


def template_metadata():
    return {"template_id":CONFIG["template_id"], "template_version":CONFIG["version"],
            "template_sha256":hashlib.sha256(_CONFIG_PATH.read_bytes()).hexdigest(), **_font_choice()}


def _contact_link(contact):
    if contact.kind == "email":
        value = contact.text.removeprefix("mailto:")
        return "mailto:" + value if "@" in value and not any(c.isspace() for c in value) else None
    if contact.kind == "link":
        value = contact.text
        return value if urlparse(value).scheme in ("https", "http") and urlparse(value).netloc else None
    return None


def _headline(ir):
    lines = [line for line in iter_lines(ir) if line.layout == "headline"]
    if len(lines) > 1:
        raise ValueError("Use at most one resume headline")
    if lines and not any(lines[0] is line for section in ir.sections for line in section.lines):
        raise ValueError("Resume headline must be a section line")
    return lines[0] if lines else None


def _kind(line, default):
    return line.layout if line.layout != "auto" else default


def _break_index(entry):
    if entry.break_before_line_id is None:
        return None
    matches = [i for i,line in enumerate(entry.bullet_lines) if line.id == entry.break_before_line_id]
    if len(matches) != 1 or matches[0] == 0:
        raise ValueError("Resume page break must identify a later existing content line")
    index = matches[0]
    if entry.bullet_lines[index].kind != "bullet" or entry.bullet_lines[index].layout == "topic":
        raise ValueError("Resume page break must precede a whole bullet")
    topic = max((i for i,line in enumerate(entry.bullet_lines[:index]) if line.layout == "topic"), default=-1)
    if topic >= 0 and not any(line.kind == "bullet" and line.layout != "topic" for line in entry.bullet_lines[topic+1:index]):
        raise ValueError("Resume page break would strand a topic heading")
    return index


def _pdf_styles():
    return {key:ParagraphStyle("Resume "+key, fontName="RolewardStandard-Bold" if spec.get("bold") else "RolewardStandard",
        fontSize=spec["size"], leading=spec["leading"], textColor=HexColor("#"+spec["color"]),
        spaceBefore=spec.get("before",0), spaceAfter=spec.get("after",0), leftIndent=spec.get("indent",0),
        bulletIndent=spec.get("bullet_indent",0), allowWidows=0, allowOrphans=0,
        bulletFontName="RolewardStandard", bulletFontSize=spec["size"],
        keepWithNext=key in ("name","headline","company","role","date","topic","degree","section"))
        for key,spec in CONFIG["styles"].items()}


class _Section(Flowable):
    def __init__(self, text, style):
        super().__init__(); self.paragraph=Paragraph(escape(text.upper()), style)
        self.spaceBefore=style.spaceBefore; self.spaceAfter=style.spaceAfter; self.keepWithNext=True
    def wrap(self,w,h):
        self.width=w; _,height=self.paragraph.wrap(w,h); self.height=height+6; return w,self.height
    def draw(self):
        self.paragraph.drawOn(self.canv,0,6)
        self.canv.setStrokeColor(HexColor("#"+CONFIG["colors"]["rule"])); self.canv.setLineWidth(.5); self.canv.line(0,0,self.width,0)


class _Row(Flowable):
    def __init__(self, left, right, style, date_style):
        super().__init__(); self.left=Paragraph(escape(left),style); self.right=Paragraph(escape(right),date_style)
        self.right.style.alignment=2; self.keepWithNext=True; self.spaceBefore=style.spaceBefore; self.spaceAfter=style.spaceAfter
    def wrap(self,w,h):
        self.width=w; self.left_width=w-126
        _,a=self.left.wrap(self.left_width,h); _,b=self.right.wrap(116,h); self.height=max(a,b); return w,self.height
    def draw(self):
        self.left.drawOn(self.canv,0,self.height-self.left.height)
        self.right.drawOn(self.canv,self.width-116,self.height-self.right.height)


class _NumberedCanvas(canvas.Canvas):
    def __init__(self,*args,**kwargs):
        kwargs.update(invariant=1,pageCompression=1); super().__init__(*args,**kwargs); self.saved=[]
    def showPage(self):
        self.saved.append(dict(self.__dict__)); self._startPage()
    def save(self):
        total=len(self.saved)
        for state in self.saved:
            self.__dict__.update(state)
            self.setFont("RolewardStandard",8); self.setFillColor(HexColor("#"+CONFIG["colors"]["secondary"]))
            self.drawString(MARGINS["left"],20,self._resume_name)
            self.drawRightString(PAGE[0]-MARGINS["right"],20,f"{self._pageNumber} / {total}")
            if self._pageNumber>1:
                self.drawString(MARGINS["left"],PAGE[1]-23,self._resume_name)
                self.drawRightString(PAGE[0]-MARGINS["right"],PAGE[1]-23,"RESUME")
            super().showPage()
        super().save()


def _entry_headers(entry, emit, row, continued=False):
    date = next((line for line in entry.header_lines if line.layout == "date"), None)
    paired = date is not None and entry.header_lines and entry.header_lines[0] is not date
    for i,line in enumerate(entry.header_lines):
        if paired and line is date:
            continue
        kind = _kind(line,"company" if i==0 else "role")
        text = line.text + (" · continued" if continued and kind == "role" else "")
        if i==0 and paired:
            row(text,date.text,kind)
        else:
            emit(text,kind)


def _skill_html(text):
    label, sep, rest = text.partition(":")
    return "<b>"+escape(label+sep)+"</b>"+escape(rest) if sep else escape(text)


def render_pdf(resume_ir: ResumeIR) -> bytes:
    _font_choice(); styles=_pdf_styles(); buffer=BytesIO()
    story=[Paragraph(escape(resume_ir.header.name),styles["name"])]
    headline=_headline(resume_ir)
    if headline: story.append(Paragraph(escape(headline.text),styles["headline"]))
    contacts=[]
    for contact in resume_ir.header.contacts:
        text=escape(contact.text); target=_contact_link(contact)
        if target: text='<link color="#'+CONFIG["colors"]["accent"]+'" href="'+escape(target,{'"':'&quot;'})+'">'+text+'</link>'
        contacts.append(text)
    story.append(Paragraph("   ·   ".join(contacts),styles["contacts"]))
    def emit(text,kind):
        content=_skill_html(text) if kind=="skill" else escape(text)
        story.append(Paragraph(content,styles[kind],bulletText="•" if kind=="bullet" else None))
    def row(left,right,kind): story.append(_Row(left,right,styles[kind],styles["date"].clone("aligned date")))
    for section in resume_ir.sections:
        story.append(_Section(section.title,styles["section"]))
        for line in section.lines:
            if line is not headline:
                emit(line.text,_kind(line,"skill" if section.kind=="skills" else "summary" if section.kind=="summary" else "body"))
        for entry in section.entries:
            _entry_headers(entry,emit,row); boundary=_break_index(entry); topic=None
            for i,line in enumerate(entry.bullet_lines):
                if i==boundary:
                    story.append(PageBreak()); _entry_headers(entry,emit,row,True)
                    if topic: emit(topic.text+" · continued","topic")
                kind=_kind(line,"bullet" if line.kind=="bullet" else "body")
                emit(line.text,kind)
                if kind=="topic": topic=line
    def make_canvas(*a,**kw):
        result=_NumberedCanvas(*a,**kw); result._resume_name=resume_ir.header.name; return result
    document=BaseDocTemplate(buffer,pagesize=PAGE,title=f"{resume_ir.header.name} Resume",author="Roleward")
    frame=Frame(MARGINS["left"],MARGINS["bottom"],WIDTH,PAGE[1]-MARGINS["top"]-MARGINS["bottom"],
        leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)
    document.addPageTemplates(PageTemplate("standard",frames=[frame]))
    document.build(story,canvasmaker=make_canvas)
    return buffer.getvalue()


def _normalized_docx(document):
    raw=BytesIO(); document.save(raw); output=BytesIO()
    with ZipFile(BytesIO(raw.getvalue())) as source, ZipFile(output,"w",ZIP_DEFLATED,compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info=ZipInfo(name,date_time=(1980,1,1,0,0,0)); info.compress_type=ZIP_DEFLATED; info.external_attr=0o600<<16
            target.writestr(info,source.read(name))
    return output.getvalue()


def render_docx(resume_ir: ResumeIR) -> bytes:
    family=_font_choice()["font_family"]; doc=Document(); sec=doc.sections[0]
    sec.page_width,sec.page_height=map(Pt,PAGE)
    for name,value in MARGINS.items(): setattr(sec,name+"_margin",Pt(value))
    sec.header_distance=sec.footer_distance=Pt(15)
    for base in doc.styles:
        if base.type==1:
            base.font.name=family; base.font.color.rgb=RGBColor(0,0,0)
            ppr=base.element.find(qn("w:pPr"))
            if ppr is not None:
                for border in list(ppr.findall(qn("w:pBdr"))): ppr.remove(border)
    for key,spec in CONFIG["styles"].items():
        s=doc.styles.add_style("Resume "+key,1); s.base_style=doc.styles["Title" if key=="name" else "Normal"]
        s.font.name=family; s.font.size=Pt(spec["size"]); s.font.bold=spec.get("bold",False); s.font.color.rgb=RGBColor.from_string(spec["color"])
        pf=s.paragraph_format; pf.space_before=Pt(spec.get("before",0)); pf.space_after=Pt(spec.get("after",0)); pf.line_spacing=Pt(spec["leading"])
        pf.left_indent=Pt(spec.get("indent",0)); pf.keep_together=True; pf.widow_control=True
        pf.keep_with_next=key in ("name","headline","company","role","date","topic","degree","section")
    doc.styles["Resume bullet"].paragraph_format.first_line_indent=Pt(-11)
    def emit(text,kind):
        p=doc.add_paragraph(style="Resume "+kind)
        if kind=="bullet":
            p.paragraph_format.tab_stops.clear_all(); p.paragraph_format.tab_stops.add_tab_stop(Pt(21)); p.add_run("•\t"+text)
        elif kind=="skill" and ":" in text:
            label,rest=text.split(":",1); p.add_run(label+":").bold=True; p.add_run(rest)
        else: p.add_run(text)
        return p
    def row(left,right,kind):
        p=emit(left,kind); p.paragraph_format.tab_stops.clear_all(); p.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH),WD_TAB_ALIGNMENT.RIGHT)
        r=p.add_run("\t"+right); r.font.size=Pt(CONFIG["styles"]["date"]["size"]); r.bold=False; r.font.color.rgb=RGBColor.from_string(CONFIG["colors"]["secondary"])
    emit(resume_ir.header.name,"name"); headline=_headline(resume_ir)
    if headline: emit(headline.text,"headline")
    contacts=emit("","contacts")
    for i,contact in enumerate(resume_ir.header.contacts):
        if i: contacts.add_run("   ·   ")
        target=_contact_link(contact)
        if not target: contacts.add_run(contact.text); continue
        h=OxmlElement("w:hyperlink"); h.set(qn("r:id"),contacts.part.relate_to(target,RT.HYPERLINK,is_external=True))
        r=OxmlElement("w:r"); pr=OxmlElement("w:rPr"); fonts=OxmlElement("w:rFonts")
        for key in ("ascii","hAnsi"): fonts.set(qn("w:"+key),family)
        pr.append(fonts)
        for tag,value in (("sz","18"),("color",CONFIG["colors"]["accent"])):
            el=OxmlElement("w:"+tag); el.set(qn("w:val"),value); pr.append(el)
        r.append(pr); t=OxmlElement("w:t"); t.text=contact.text; r.append(t); h.append(r); contacts._p.append(h)
    for section in resume_ir.sections:
        p=emit(section.title.upper(),"section")
        for run in p.runs:
            el=OxmlElement("w:spacing"); el.set(qn("w:val"),"20"); run._r.get_or_add_rPr().append(el)
        borders=OxmlElement("w:pBdr"); bottom=OxmlElement("w:bottom")
        for key,value in {"val":"single","sz":"4","space":"4","color":CONFIG["colors"]["rule"]}.items(): bottom.set(qn("w:"+key),value)
        borders.append(bottom); p._p.get_or_add_pPr().append(borders)
        for line in section.lines:
            if line is not headline: emit(line.text,_kind(line,"skill" if section.kind=="skills" else "summary" if section.kind=="summary" else "body"))
        for entry in section.entries:
            _entry_headers(entry,emit,row); boundary=_break_index(entry); topic=None
            for i,line in enumerate(entry.bullet_lines):
                if i==boundary:
                    doc.add_page_break(); _entry_headers(entry,emit,row,True)
                    if topic: emit(topic.text+" · continued","topic")
                kind=_kind(line,"bullet" if line.kind=="bullet" else "body"); emit(line.text,kind)
                if kind=="topic": topic=line
    sec.different_first_page_header_footer=True
    p=sec.header.paragraphs[0]; p.style=doc.styles["Resume navigation"]
    p.paragraph_format.tab_stops.clear_all(); p.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH),WD_TAB_ALIGNMENT.RIGHT); p.add_run(resume_ir.header.name+"\tRESUME")
    for footer in (sec.footer,sec.first_page_footer):
        p=footer.paragraphs[0]; p.style=doc.styles["Resume navigation"]
        p.paragraph_format.tab_stops.clear_all(); p.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH),WD_TAB_ALIGNMENT.RIGHT); p.add_run(resume_ir.header.name+"\t")
        for instruction in ("PAGE","NUMPAGES"):
            if instruction=="NUMPAGES": p.add_run(" / ")
            field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),instruction); p._p.append(field)
    fixed_time=datetime(2000,1,1,tzinfo=timezone.utc)
    props=doc.core_properties; props.author=props.last_modified_by="Roleward"; props.title=f"{resume_ir.header.name} Resume"
    props.subject="Standard tailored resume"; props.created=props.modified=fixed_time; props.revision=1
    return _normalized_docx(doc)


def render_resume(resume_ir: ResumeIR, export_format: str) -> tuple[bytes,str]:
    if export_format=="pdf": return render_pdf(resume_ir),"application/pdf"
    if export_format=="docx": return render_docx(resume_ir),"application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    raise ValueError("Unsupported Resume export format.")
