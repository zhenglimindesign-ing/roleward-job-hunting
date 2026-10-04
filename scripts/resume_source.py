#!/usr/bin/env python3
"""Extract a readable resume source without treating its claims as confirmed truth."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path


def extract_source(path: Path) -> dict:
    data = path.read_bytes()
    if len(data) > 4 * 1024 * 1024:
        raise ValueError("Resume source exceeds the supported 4 MB limit")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from io import BytesIO
        from pdfminer.high_level import extract_text
        from pdfminer.pdfpage import PDFPage
        if len(list(PDFPage.get_pages(BytesIO(data), check_extractable=True))) > 20:
            raise ValueError("Resume source exceeds the supported 20-page limit")
        pages = extract_text(BytesIO(data)).split("\f")
    elif suffix == ".docx":
        from io import BytesIO
        from docx import Document
        document = Document(BytesIO(data))
        if document.tables:
            raise ValueError("Table-based DOCX requires host-assisted extraction and coverage review")
        pages = ["\n".join(p.text for p in document.paragraphs)]
    elif suffix in (".txt", ".md"):
        pages = [data.decode("utf-8")]
    else:
        raise ValueError("Supported sources: readable PDF, paragraph-based DOCX, UTF-8 TXT or Markdown")
    lines = []
    for page, text in enumerate(pages, 1):
        for line in text.splitlines():
            if line.strip():
                lines.append({"id": f"source-{len(lines) + 1}", "text": line.strip(), "page": page})
    if not lines:
        raise ValueError("No readable text layer; OCR is outside this helper's supported path")
    return {"source_sha256": hashlib.sha256(data).hexdigest(), "parser_version": "source-text-v1",
            "authority": "source_material", "lines": lines,
            "coverage_review_required": True}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = extract_source(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
